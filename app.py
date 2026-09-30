import sqlite3
from datetime import date
from io import BytesIO
from pathlib import Path
import pandas as pd
import streamlit as st
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak, Image as RLImage
from reportlab.lib.utils import ImageReader


APP_DIR=Path(__file__).resolve().parent; DB=APP_DIR/'stock_tracker.db'
st.set_page_config(page_title='Stock Profit Tracker Pro',page_icon='📈',layout='wide')
st.markdown("""
<style>
/* Responsive layout for phone screens */
@media (max-width: 640px) {
    /* Main content */
    [data-testid="stMainBlockContainer"] {
        padding-left: 0.65rem !important;
        padding-right: 0.65rem !important;
    }

    /* Keep Streamlit tabs in a single non-overlapping row */
    div[data-testid="stTabs"] {
        width: 100% !important;
        min-width: 0 !important;
    }

    div[data-testid="stTabs"] > div:first-child {
        width: 100% !important;
        min-width: 0 !important;
        overflow-x: auto !important;
        overflow-y: hidden !important;
        scrollbar-width: none !important;
        -webkit-overflow-scrolling: touch !important;
    }

    div[data-testid="stTabs"] > div:first-child::-webkit-scrollbar {
        display: none !important;
    }

    div[data-testid="stTabs"] [role="tablist"] {
        display: flex !important;
        flex-wrap: nowrap !important;
        width: max-content !important;
        min-width: 100% !important;
        gap: 0 !important;
    }

    div[data-testid="stTabs"] [role="tab"] {
        flex: 0 0 auto !important;
        width: auto !important;
        min-width: 0 !important;
        white-space: nowrap !important;
        padding: 0.45rem 0.55rem !important;
        margin: 0 !important;
        font-size: 0.82rem !important;
    }

    /* Prevent metric cards from forcing horizontal overflow */
    div[data-testid="stMetric"] {
        min-width: 0 !important;
        overflow: hidden !important;
    }

    div[data-testid="stMetricValue"] {
        font-size: 1.35rem !important;
        white-space: nowrap !important;
        overflow: hidden !important;
        text-overflow: ellipsis !important;
    }

    /* Tables can scroll locally instead of widening the whole page */
    [data-testid="stDataFrame"] {
        max-width: 100% !important;
        overflow-x: auto !important;
    }
}
</style>
""", unsafe_allow_html=True)


def conn():
    c=sqlite3.connect(DB); c.row_factory=sqlite3.Row; return c

def init():
    c=conn(); c.execute('''CREATE TABLE IF NOT EXISTS transactions(id INTEGER PRIMARY KEY AUTOINCREMENT,person TEXT NOT NULL,stock_name TEXT NOT NULL,sell_date TEXT NOT NULL,quantity INTEGER NOT NULL CHECK(quantity>0),buying_price REAL NOT NULL CHECK(buying_price>=0),selling_price REAL NOT NULL CHECK(selling_price>=0),created_at TEXT DEFAULT CURRENT_TIMESTAMP)''')
    if c.execute('SELECT COUNT(*) FROM transactions').fetchone()[0]==0:
        c.executemany('INSERT INTO transactions(person,stock_name,sell_date,quantity,buying_price,selling_price) VALUES(?,?,?,?,?,?)',[
            ('Papa','HFCL','2026-05-13',1000,17.4,147.4),('Papa','Indo MIM','2026-07-30',120,485,707),('Papa','Indo MIM','2026-07-30',100,485,719),('Papa','Indo MIM','2026-09-22',100,485,1173.4),('Papa','Glasswall Systems','2026-09-22',648,182,307.5)])
    c.commit(); c.close()

def getdf():
    c=conn(); rows=c.execute('SELECT * FROM transactions ORDER BY sell_date DESC,id DESC').fetchall(); c.close()
    a=[]
    for r in rows:
        p=round(r['selling_price']-r['buying_price'],2); t=round(p*r['quantity'],2)
        a.append({'ID':r['id'],'Name':r['person'],'Stock Name':r['stock_name'],'Sell Date':pd.to_datetime(r['sell_date']).date(),'Quantity':r['quantity'],'Buying Price':r['buying_price'],'Selling Price':r['selling_price'],'Profit per Share':p,'Total Profit':t})
    return pd.DataFrame(a)

def add(person,stock,d,q,b,s):
    c=conn(); c.execute('INSERT INTO transactions(person,stock_name,sell_date,quantity,buying_price,selling_price) VALUES(?,?,?,?,?,?)',(person,stock.strip(),d.isoformat(),int(q),float(b),float(s))); c.commit(); c.close()

def delete(i):
    c=conn(); c.execute('DELETE FROM transactions WHERE id=?',(int(i),)); c.commit(); c.close()

def clear():
    c=conn(); c.execute('DELETE FROM transactions'); c.commit(); c.close()

def money(x): return f'₹{x:,.0f}'

def excel(df):
    out=BytesIO(); x=df.drop(columns=['ID']).copy(); x['Sell Date']=pd.to_datetime(x['Sell Date']).dt.date
    with pd.ExcelWriter(out,engine='openpyxl') as w:
        x.to_excel(w,index=False,sheet_name='Transactions'); wb=w.book; ws=w.sheets['Transactions']
        fill=PatternFill('solid',fgColor='1F4E78'); font=Font(color='FFFFFF',bold=True); side=Side(style='thin',color='D9E1F2')
        for cell in ws[1]: cell.fill=fill; cell.font=font; cell.alignment=Alignment(horizontal='center'); cell.border=Border(bottom=side)
        for row in ws.iter_rows(min_row=2):
            row[2].number_format='dd mmm yyyy'; row[3].number_format='#,##0'
            for j in [4,5,6,7]: row[j].number_format='₹#,##0'
        tr=ws.max_row+2
        ws.cell(tr,1,'TOTAL')
        ws.cell(tr,4,int(x['Quantity'].sum()))
        ws.cell(tr,5,float((x['Buying Price']*x['Quantity']).sum()))
        ws.cell(tr,6,float((x['Selling Price']*x['Quantity']).sum()))
        ws.cell(tr,7,'')
        ws.cell(tr,8,float(x['Total Profit'].sum()))
        for col in [5,6,8]:
            ws.cell(tr,col).number_format='₹#,##0'
        for c in ws[tr]:
            c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='E2F0D9')
        for i,wid in enumerate([14,24,16,12,17,17,19,17],1): ws.column_dimensions[get_column_letter(i)].width=wid
        ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:H{tr-2}'
        s=wb.create_sheet('Summary'); inv=float((x['Buying Price']*x['Quantity']).sum()); sales=float((x['Selling Price']*x['Quantity']).sum()); profit=float(x['Total Profit'].sum())
        vals=[['Stock Profit Tracker — Summary',''],['Transactions',len(x)],['Total Quantity',int(x['Quantity'].sum())],['Total Investment',inv],['Total Sales Value',sales],['Total Profit',profit]]
        for r in vals:s.append(r)
        for c in s[1]:c.fill=fill;c.font=Font(color='FFFFFF',bold=True,size=14)
        for r in range(4,s.max_row+1):
            s.cell(r,2).number_format='₹#,##0'
        s.column_dimensions['A'].width=26;s.column_dimensions['B'].width=22
    return out.getvalue()

def pdf(df):
    """Create a PDF containing KPI summary, analytics tables/charts, and transactions."""
    import tempfile
    import os
    import matplotlib.pyplot as plt

    out=BytesIO()
    doc=SimpleDocTemplate(
        out,
        pagesize=landscape(A4),
        rightMargin=8*mm,leftMargin=8*mm,topMargin=8*mm,bottomMargin=8*mm
    )
    stl=getSampleStyleSheet()
    title=ParagraphStyle('t',parent=stl['Title'],alignment=TA_CENTER,fontSize=18)
    h2=ParagraphStyle('h2',parent=stl['Heading2'],fontSize=13,spaceBefore=5*mm,spaceAfter=2*mm)
    small=ParagraphStyle('small',parent=stl['BodyText'],fontSize=8)
    # ReportLab's built-in Helvetica does not contain the Indian rupee glyph.
    # Use "Rs." in the PDF to avoid missing-glyph black boxes.
    def pdf_money(x):
        return f"Rs. {x:,.0f}"

    inv=float((df['Buying Price']*df['Quantity']).sum()) if len(df) else 0
    sales=float((df['Selling Price']*df['Quantity']).sum()) if len(df) else 0
    profit=float(df['Total Profit'].sum()) if len(df) else 0
    qty=int(df['Quantity'].sum()) if len(df) else 0
    ret=(profit/inv*100) if inv else 0
    profitable=int((df['Total Profit']>0).sum())
    loss=int((df['Total Profit']<0).sum())
    breakeven=int((df['Total Profit']==0).sum())

    story=[
        Paragraph('Stock Profit Tracker — Profit & Analytics Report',title),
        Paragraph(f'Report generated on: {date.today().strftime("%d %b %Y")}',small),
        Spacer(1,4*mm)
    ]

    # KPI summary
    sm=Table([
        ['Transactions','Quantity','Investment','Sales Value','Total Profit','Return'],
        [str(len(df)),f"{qty:,}",pdf_money(inv),pdf_money(sales),pdf_money(profit),f"{ret:.2f}%"]
    ],colWidths=[30*mm,30*mm,45*mm,45*mm,45*mm,30*mm])
    sm.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#1F4E78')),
        ('TEXTCOLOR',(0,0),(-1,0),colors.white),
        ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),
        ('BACKGROUND',(0,1),(-1,1),colors.HexColor('#E2F0D9')),
        ('GRID',(0,0),(-1,-1),.4,colors.grey),
        ('ALIGN',(0,0),(-1,-1),'CENTER'),
        ('FONTSIZE',(0,0),(-1,-1),9)
    ]))
    story += [sm, Spacer(1,3*mm)]

    status=Table([
        ['Profitable Transactions','Loss Transactions','Break-even Transactions'],
        [str(profitable),str(loss),str(breakeven)]
    ],colWidths=[55*mm,55*mm,55*mm])
    status.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#D9EAF7')),
        ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),
        ('GRID',(0,0),(-1,-1),.4,colors.grey),
        ('ALIGN',(0,0),(-1,-1),'CENTER')
    ]))
    story += [status]

    # Detailed transactions FIRST
    story.append(Paragraph('Detailed Transactions',h2))
    rows=[['Name','Stock','Sell Date','Qty','Buy','Sell','Profit/Share','Total Profit']]
    for _,r in df.iterrows():
        rows.append([
            r['Name'],r['Stock Name'],
            pd.to_datetime(r['Sell Date']).strftime('%d %b %Y'),
            f"{int(r['Quantity']):,}",
            pdf_money(r['Buying Price']),pdf_money(r['Selling Price']),
            pdf_money(r['Profit per Share']),pdf_money(r['Total Profit'])
        ])
    tb=Table(rows,repeatRows=1,colWidths=[23*mm,42*mm,29*mm,18*mm,28*mm,28*mm,35*mm,38*mm])
    tb.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#1F4E78')),
        ('TEXTCOLOR',(0,0),(-1,0),colors.white),
        ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),
        ('GRID',(0,0),(-1,-1),.3,colors.grey),
        ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F5F8FA')]),
        ('FONTSIZE',(0,0),(-1,-1),8),
        ('ALIGN',(2,1),(-1,-1),'RIGHT')
    ]))
    total_row=Table(
        [['TOTAL','','',f"{qty:,}",'', '', '',pdf_money(profit)]],
        colWidths=[23*mm,42*mm,29*mm,18*mm,28*mm,28*mm,35*mm,38*mm]
    )
    total_row.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#E2F0D9')),
        ('FONTNAME',(0,0),(-1,-1),'Helvetica-Bold'),
        ('GRID',(0,0),(-1,-1),.4,colors.HexColor('#7F8C8D')),
        ('ALIGN',(3,0),(3,0),'RIGHT'),
        ('ALIGN',(7,0),(7,0),'RIGHT'),
        ('FONTSIZE',(0,0),(-1,-1),8)
    ]))
    story += [
        tb,
        total_row,
        Spacer(1,5*mm)
    ]

    story.append(PageBreak())

    # Analytics tables
    person=df.groupby('Name',as_index=False).agg(
        Transactions=('ID','count'),
        Quantity=('Quantity','sum'),
        Investment=('Buying Price',lambda x: 0)
    )
    # Recalculate investment/sales at row level so grouped values are accurate.
    tmp=df.copy()
    tmp['Investment']=tmp['Buying Price']*tmp['Quantity']
    tmp['Sales Value']=tmp['Selling Price']*tmp['Quantity']
    person=tmp.groupby('Name',as_index=False).agg(
        Transactions=('ID','count'),
        Quantity=('Quantity','sum'),
        Investment=('Investment','sum'),
        Sales=('Sales Value','sum'),
        Profit=('Total Profit','sum')
    )

    stock=tmp.groupby('Stock Name',as_index=False).agg(
        Transactions=('ID','count'),
        Quantity=('Quantity','sum'),
        Investment=('Investment','sum'),
        Sales=('Sales Value','sum'),
        Profit=('Total Profit','sum')
    )

    tmp['Month']=pd.to_datetime(tmp['Sell Date']).dt.to_period('M').astype(str)
    monthly=tmp.groupby('Month',as_index=False).agg(
        Transactions=('ID','count'),
        Quantity=('Quantity','sum'),
        Investment=('Investment','sum'),
        Sales=('Sales Value','sum'),
        Profit=('Total Profit','sum')
    )

    def analytics_table(frame, first_col, widths):
        headers=[first_col,'Transactions','Quantity','Investment','Sales','Profit']
        rows=[headers]
        for _,r in frame.iterrows():
            rows.append([
                str(r[first_col]),
                f"{int(r['Transactions']):,}",
                f"{int(r['Quantity']):,}",
                pdf_money(r['Investment']),
                pdf_money(r['Sales']),
                pdf_money(r['Profit'])
            ])
        tb=Table(rows,colWidths=widths,repeatRows=1)
        tb.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#1F4E78')),
            ('TEXTCOLOR',(0,0),(-1,0),colors.white),
            ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),
            ('GRID',(0,0),(-1,-1),.3,colors.grey),
            ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F5F8FA')]),
            ('FONTSIZE',(0,0),(-1,-1),8),
            ('ALIGN',(1,1),(-1,-1),'RIGHT')
        ]))
        return tb

    story += [
        Paragraph('Analytics — Person-wise',h2),
        analytics_table(person,'Name',[40*mm,30*mm,30*mm,45*mm,45*mm,45*mm]),
        Paragraph('Analytics — Stock-wise',h2),
        analytics_table(stock,'Stock Name',[40*mm,30*mm,30*mm,45*mm,45*mm,45*mm]),
        Paragraph('Analytics — Monthly',h2),
        analytics_table(monthly,'Month',[40*mm,30*mm,30*mm,45*mm,45*mm,45*mm])
    ]

    # Generate analytics charts temporarily and embed them in the PDF.
    with tempfile.TemporaryDirectory() as td:
        chart_paths=[]

        def save_bar(labels, values, title_text, filename, ylabel='Profit (Rs.)'):
            if len(labels)==0:
                return None
            fig,ax=plt.subplots(figsize=(8.5,3.2))
            ax.bar(labels, values)
            ax.set_title(title_text)
            ax.set_ylabel(ylabel)
            ax.tick_params(axis='x',rotation=35)
            fig.tight_layout()
            path=os.path.join(td,filename)
            fig.savefig(path,dpi=150,bbox_inches='tight')
            plt.close(fig)
            return path

        def save_line(labels, values, title_text, filename):
            if len(labels)==0:
                return None
            fig,ax=plt.subplots(figsize=(8.5,3.2))
            ax.plot(labels, values, marker='o')
            ax.set_title(title_text)
            ax.set_ylabel('Profit (Rs.)')
            ax.tick_params(axis='x',rotation=35)
            fig.tight_layout()
            path=os.path.join(td,filename)
            fig.savefig(path,dpi=150,bbox_inches='tight')
            plt.close(fig)
            return path

        p=save_bar(person['Name'].tolist(),person['Profit'].tolist(),'Profit by Person','profit_by_person.png')
        if p: chart_paths.append(p)
        p=save_bar(stock['Stock Name'].tolist(),stock['Profit'].tolist(),'Profit by Stock','profit_by_stock.png')
        if p: chart_paths.append(p)
        p=save_line(monthly['Month'].tolist(),monthly['Profit'].tolist(),'Monthly Profit Trend','monthly_profit.png')
        if p: chart_paths.append(p)

        story.append(PageBreak())
        story.append(Paragraph('Analytics Charts',h2))
        for cp in chart_paths:
            story.append(RLImage(cp,width=250*mm,height=90*mm))
            story.append(Spacer(1,3*mm))

        doc.build(story)

    return out.getvalue()


init(); df=getdf()

# Admin access: no signup/login system. Visitors remain read-only unless
# they know the admin password stored privately in Streamlit Secrets.
try:
    ADMIN_PASSWORD=st.secrets["ADMIN_PASSWORD"]
except Exception:
    ADMIN_PASSWORD=""

st.title('📈 Stock Profit Tracker Pro')
st.caption('Public read-only stock analytics dashboard. Transaction changes are restricted to Admin mode.')

with st.sidebar:
    st.header('🔐 Admin Access')
    if ADMIN_PASSWORD:
        entered_password=st.text_input('Admin Password',type='password',placeholder='Enter admin password')
        is_admin=bool(entered_password) and entered_password==ADMIN_PASSWORD
        if is_admin:
            st.success('Admin mode enabled.')
        elif entered_password:
            st.error('Incorrect admin password.')
        else:
            st.info('Viewer mode — transaction data cannot be changed.')
    else:
        is_admin=False
        st.warning('Admin access is not configured. The app is read-only.')

    st.divider()

    if is_admin:
        st.header('➕ Add Transaction')
        with st.form('add',clear_on_submit=True):
            person=st.text_input('Name',placeholder='Enter any name — Papa, Mummy, Naman, Rahul, etc.')
            stock=st.text_input('Stock Name')
            d=st.date_input('Sell Date',date.today())
            q=st.number_input('Quantity',1,10000000,1,1)
            b=st.number_input('Buying Price (₹)',0.0,100000000.0,0.0,.10)
            s=st.number_input('Selling Price (₹)',0.0,100000000.0,0.0,.10)
            ok=st.form_submit_button('Add Transaction',type='primary',use_container_width=True)
            if ok:
                if not stock.strip():
                    st.error('Stock name is required.')
                else:
                    add(person,stock,d,q,b,s)
                    st.success(f'Added. Profit/share: {money(s-b)} | Total: {money((s-b)*q)}')
                    st.rerun()

        st.divider()
        st.header('📥 Import')
        up=st.file_uploader('Excel/CSV',type=['xlsx','csv'])
        if up and st.button('Import Rows',use_container_width=True):
            try:
                imp=pd.read_csv(up) if up.name.lower().endswith('.csv') else pd.read_excel(up)
                req=['Name','Stock Name','Sell Date','Quantity','Buying Price','Selling Price']
                miss=[x for x in req if x not in imp.columns]
                if miss:
                    st.error('Missing: '+', '.join(miss))
                else:
                    c=conn()
                    for _,r in imp.iterrows():
                        c.execute(
                            'INSERT INTO transactions(person,stock_name,sell_date,quantity,buying_price,selling_price) VALUES(?,?,?,?,?,?)',
                            (str(r['Name']),str(r['Stock Name']),
                             pd.to_datetime(r['Sell Date']).date().isoformat(),
                             int(float(r['Quantity'])),float(r['Buying Price']),float(r['Selling Price']))
                        )
                    c.commit(); c.close()
                    st.success(f'Imported {len(imp)} rows.')
                    st.rerun()
            except Exception as e:
                st.error(str(e))

        st.divider()
        st.header('🧹 Data Management')
        if st.button('Delete All Transactions',use_container_width=True):
            st.session_state.confirm=True
        if st.session_state.get('confirm'):
            st.warning('This permanently deletes the transaction database.')
            a,bx=st.columns(2)
            if a.button('Confirm'):
                clear()
                st.session_state.confirm=False
                st.rerun()
            if bx.button('Cancel'):
                st.session_state.confirm=False
                st.rerun()
    else:
        st.header('👁️ Viewer Mode')
        st.caption('You can view analytics and download reports. Only Admin can add, import, or delete transactions.')

# Dashboard filters — defaults show the complete dataset.
with st.expander('🎛️ Dashboard Filters', expanded=False):
    view_df=df.copy()
    if len(df):
        f1,f2,f3=st.columns(3)
        people=['All']+sorted(df['Name'].dropna().astype(str).unique().tolist())
        stocks=['All']+sorted(df['Stock Name'].dropna().astype(str).unique().tolist())
        selected_person=f1.selectbox('Person',people)
        selected_stock=f2.selectbox('Stock',stocks)
        min_date=pd.to_datetime(df['Sell Date']).min().date()
        max_date=pd.to_datetime(df['Sell Date']).max().date()
        date_range=f3.date_input('Sell Date Range',(min_date,max_date),min_value=min_date,max_value=max_date)
        if not isinstance(date_range,tuple) or len(date_range)!=2:
            date_range=(min_date,max_date)
        if selected_person!='All':
            view_df=view_df[view_df['Name']==selected_person]
        if selected_stock!='All':
            view_df=view_df[view_df['Stock Name']==selected_stock]
        view_df=view_df[
            (pd.to_datetime(view_df['Sell Date']).dt.date>=date_range[0]) &
            (pd.to_datetime(view_df['Sell Date']).dt.date<=date_range[1])
        ].copy()
        if len(view_df)!=len(df):
            st.caption(f'Showing {len(view_df)} of {len(df)} transactions based on the selected filters.')
    else:
        st.info('No transactions available to filter.')

profit=float(view_df['Total Profit'].sum()) if len(view_df) else 0
inv=float((view_df['Buying Price']*view_df['Quantity']).sum()) if len(view_df) else 0
sales=float((view_df['Selling Price']*view_df['Quantity']).sum()) if len(view_df) else 0
qty=int(view_df['Quantity'].sum()) if len(view_df) else 0
transactions=len(view_df)
profitable_count=int((view_df['Total Profit']>0).sum()) if len(view_df) else 0
profit_rate=(profitable_count/transactions*100) if transactions else 0
avg_profit=profit/transactions if transactions else 0
margin=profit/sales*100 if sales else 0

k=st.columns(5)
with k[0]:
    st.markdown(f'''
    <div style="border:2px solid #2E7D32;border-radius:14px;padding:12px 14px;
    background:linear-gradient(135deg,#EAF7EC,#F5FBF6);box-shadow:0 2px 8px rgba(46,125,50,.14);min-height:92px;">
      <div style="font-size:.82rem;font-weight:700;color:#1B5E20;margin-bottom:5px;">💰 TOTAL PROFIT</div>
      <div style="font-size:1.65rem;line-height:1.15;font-weight:900;color:#1B5E20;white-space:nowrap;">{money(profit)}</div>
    </div>''',unsafe_allow_html=True)
k[1].metric('Investment',money(inv))
k[2].metric('Sales Value',money(sales))
k[3].metric('Quantity',f'{qty:,}')
k[4].metric('Return',f'{profit/inv*100:.2f}%' if inv else '0.00%')

st.markdown('### 📌 Performance Snapshot')
s1,s2,s3,s4,s5=st.columns(5)
s1.metric('Transactions',f'{transactions:,}')
s2.metric('Profitable Trades',f'{profitable_count:,}')
s3.metric('Profit Rate',f'{profit_rate:.1f}%')
s4.metric('Avg. Profit / Trade',money(avg_profit))
s5.metric('Profit Margin',f'{margin:.1f}%')

if len(view_df):
    by_stock=view_df.groupby('Stock Name')['Total Profit'].sum().sort_values(ascending=False)
    by_person=view_df.groupby('Name')['Total Profit'].sum().sort_values(ascending=False)
    best_stock=str(by_stock.index[0])
    best_person=str(by_person.index[0])
    best_trade=view_df.loc[view_df['Total Profit'].idxmax()]
    latest_date=pd.to_datetime(view_df['Sell Date']).max().strftime('%d %b %Y')
    q1,q2,q3,q4=st.columns(4)
    q1.info(f'🏆 **Top Stock by Profit**\n\n{best_stock}\n\n{money(by_stock.iloc[0])}')
    q2.info(f'👤 **Top Person by Profit**\n\n{best_person}\n\n{money(by_person.iloc[0])}')
    q3.info(f'🚀 **Best Single Trade**\n\n{best_trade["Stock Name"]}\n\n{money(best_trade["Total Profit"])}')
    q4.info(f'📅 **Latest Sale Date**\n\n{latest_date}\n\n{transactions:,} transaction(s) in view')


t1,t2,t3=st.tabs(['📋 Transactions','📊 Analytics','📤 Export'])

with t1:
    if len(view_df):
        v=view_df.copy()
        v['Sell Date']=pd.to_datetime(v['Sell Date']).dt.strftime('%d %b %Y')
        v=v.drop(columns=['ID'])
        styled_v=v.style.set_properties(
            subset=['Profit per Share','Total Profit'],
            **{'font-weight':'800'}
        ).format({
            'Buying Price':'{:,.0f}',
            'Selling Price':'{:,.0f}',
            'Profit per Share':'{:,.0f}',
            'Total Profit':'{:,.0f}'
        })
        st.dataframe(styled_v,use_container_width=True,hide_index=True)

        if is_admin:
            opts={f"#{int(r['ID'])} — {r['Stock Name']} — {r['Sell Date']} — {money(r['Total Profit'])}":int(r['ID']) for _,r in view_df.iterrows()}
            sel=st.selectbox('Delete transaction',list(opts))
            if st.button('Delete Selected',type='secondary'):
                delete(opts[sel])
                st.rerun()
    else:
        st.info('No transactions yet.')

with t2:
    if len(view_df):
        a,b=st.columns(2)
        ps=view_df.groupby('Name')['Total Profit'].sum()
        ss=view_df.groupby('Stock Name')['Total Profit'].sum()
        a.subheader('Profit by Person'); a.bar_chart(ps)
        b.subheader('Profit by Stock'); b.bar_chart(ss)

        m=view_df.copy()
        m['Month']=pd.to_datetime(m['Sell Date']).dt.to_period('M').astype(str)
        st.subheader('Monthly Profit')
        st.line_chart(m.groupby('Month')['Total Profit'].sum())
        st.markdown('### 📈 Profit Contribution')
        ca,cb=st.columns(2)
        ca.bar_chart(view_df.groupby('Stock Name')['Total Profit'].sum().sort_values(ascending=False))
        cb.bar_chart(view_df.groupby('Name')['Total Profit'].sum().sort_values(ascending=False))

        st.markdown('### 📋 Stock Performance Summary')
        summary_view=view_df.copy()
        summary_view['Investment']=summary_view['Buying Price']*summary_view['Quantity']
        summary_view['Sales Value']=summary_view['Selling Price']*summary_view['Quantity']
        summary=summary_view.groupby('Stock Name',as_index=False).agg(
            Transactions=('ID','count'),Quantity=('Quantity','sum'),
            Investment=('Investment','sum'),Sales=('Sales Value','sum'),
            Profit=('Total Profit','sum')
        )
        summary['Return %']=(summary['Profit']/summary['Investment'].replace(0,pd.NA)*100).round(1)
        summary=summary.sort_values('Profit',ascending=False)
        st.dataframe(
            summary[['Stock Name','Transactions','Quantity','Investment','Sales','Profit','Return %']].style.format({
                'Investment':'{:,.0f}',
                'Sales':'{:,.0f}',
                'Profit':'{:,.0f}',
                'Return %':'{:.1f}%'
            }),
            use_container_width=True,hide_index=True
        )


        audit=view_df.copy()
        audit['Calculated Profit/Share']=(audit['Selling Price']-audit['Buying Price']).round(2)
        audit['Calculated Total Profit']=(audit['Calculated Profit/Share']*audit['Quantity']).round(2)
        audit['Profit/Share OK']=audit['Profit per Share'].round(2)==audit['Calculated Profit/Share']
        audit['Total Profit OK']=audit['Total Profit'].round(2)==audit['Calculated Total Profit']
        st.subheader('Calculation Audit')
        st.dataframe(
            audit[['Stock Name','Profit per Share','Calculated Profit/Share','Profit/Share OK',
                   'Total Profit','Calculated Total Profit','Total Profit OK']].style.format({
                'Profit per Share':'{:,.0f}',
                'Calculated Profit/Share':'{:,.0f}',
                'Total Profit':'{:,.0f}',
                'Calculated Total Profit':'{:,.0f}'
            }),
            use_container_width=True,hide_index=True
        )
        if audit['Profit/Share OK'].all() and audit['Total Profit OK'].all():
            st.success('All calculations are correct.')
        else:
            st.error('Calculation mismatch detected.')
    else:
        st.info('Add transactions to see analytics.')

with t3:
    if len(view_df):
        c1,c2=st.columns(2)
        c1.download_button(
            '⬇️ Download Excel',
            excel(view_df),
            f'Stock_Profit_Report_{date.today()}.xlsx',
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            use_container_width=True
        )
        c2.download_button(
            '⬇️ Download PDF',
            pdf(view_df),
            f'Stock_Profit_Report_{date.today()}.pdf',
            'application/pdf',
            use_container_width=True
        )
        st.download_button(
            'Download CSV Backup',
            view_df.drop(columns=['ID']).to_csv(index=False).encode(),
            f'Stock_Profit_Backup_{date.today()}.csv',
            'text/csv',
            use_container_width=True
        )
    else:
        st.info('Add transactions before exporting.')

st.divider()
st.caption('Use Dashboard Filters to focus the report by person, stock, or date range. Profit = Selling Price − Buying Price; Total Profit = Profit/Share × Quantity. No brokerage, STT, GST, taxes, or other charges are included.')
